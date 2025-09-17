from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException
from fastapi.responses import JSONResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field
from typing import List, Optional, Dict, Any
import uuid
from datetime import datetime
import io
from openpyxl import load_workbook
from exponent_server_sdk import PushClient, PushMessage, PushServerError
import asyncio

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# Create the main app without a prefix
app = FastAPI()

# Create a router with the /api prefix
api_router = APIRouter(prefix="/api")

# ----------------------------------------------------------------------------
# Models
# ----------------------------------------------------------------------------
class StatusCheck(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_name: str
    timestamp: datetime = Field(default_factory=datetime.utcnow)

class StatusCheckCreate(BaseModel):
    client_name: str

class PushTokenIn(BaseModel):
    manager: str
    token: str

# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
GOALS = [
    {"label": "7d/40h", "days": 7, "hours": 40},
    {"label": "20d/60h", "days": 20, "hours": 60},
    {"label": "22d/80h", "days": 22, "hours": 80},
]
DIAMOND_MILESTONES = [50000, 100000, 300000]

REQUIRED_HEADERS = [
    "Creator ID",
    "Creator's username",
    "Joined time",
    "Diamonds",
    "LIVE duration",
    "Valid go LIVE days",
]

OPTIONAL_HEADERS = [
    "manager",
    "Data period",
    "Days since joining",
]


def normalize_header_map(headers: List[str]) -> Dict[str, str]:
    mapping = {}
    for h in headers:
        mapping[h.strip().lower()] = h
    return mapping


def get_cell(row: Dict[str, Any], hmap: Dict[str, str], key_variants: List[str], default=None):
    for kv in key_variants:
        lk = kv.lower()
        if lk in hmap:
            return row.get(hmap[lk], default)
    return default


def compute_goal_status(valid_days: int, live_hours: float) -> Dict[str, Any]:
    statuses = []
    achieved_levels = []
    near_goal = None
    for idx, g in enumerate(GOALS):
        achieved = valid_days >= g["days"] and live_hours >= g["hours"]
        achieved_levels.append(achieved)
        statuses.append({"label": g["label"], "achieved": achieved})

    # Determine next goal proximity
    for idx, g in enumerate(GOALS):
        if not achieved_levels[idx]:
            days_left = max(g["days"] - valid_days, 0)
            hours_left = max(g["hours"] - live_hours, 0.0)
            # Near if within 3 days or 12 hours of requirement
            if days_left <= 3 or hours_left <= 12:
                near_goal = {
                    "label": g["label"],
                    "days_left": days_left,
                    "hours_left": round(hours_left, 2),
                }
            break

    # Overall status
    if all(achieved_levels):
        status = {"code": "SUPERSTAR", "icon": "trophy"}
    elif any(achieved_levels):
        status = {"code": "ACHIEVED_PARTIAL", "icon": "target"}
    elif near_goal is not None:
        status = {"code": "NEAR", "icon": "alert"}
    else:
        status = {"code": "IN_PROGRESS", "icon": "hourglass"}

    return {"levels": statuses, "near_goal": near_goal, "status": status}


def diamonds_milestones_near(diamonds: int) -> List[Dict[str, Any]]:
    res = []
    for m in DIAMOND_MILESTONES:
        # Near threshold set to 80% per business rule
        if diamonds < m and diamonds >= int(0.8 * m):
            res.append({"milestone": m, "progress": round(diamonds / m, 3)})
    return res


async def upsert_creator(doc: Dict[str, Any]):
    await db.creators.update_one(
        {"creator_id": doc["creator_id"]},
        {"$set": doc, "$currentDate": {"updated_at": True}},
        upsert=True,
    )


async def push_summary_notifications(period: str):
    # Aggregate by manager and send a small summary
    pipeline = [
        {"$group": {
            "_id": {"manager": {"$ifNull": ["$manager", "Unassigned"]}},
            "count_creators": {"$sum": 1},
            "alerts": {
                "$sum": {
                    "$cond": [
                        {"$gt": [{"$size": {"$ifNull": ["$alerts", []]}}, 0]}, 1, 0
                    ]
                }
            },
        }},
    ]
    groups = await db.creators.aggregate(pipeline).to_list(length=1000)

    for g in groups:
        manager = g["_id"]["manager"]
        tokens = await db.push_tokens.find({"manager": manager}).to_list(100)
        if not tokens:
            continue
        body = f"{period}: {g['count_creators']} creadores, {g['alerts']} con alertas"
        messages = [
            PushMessage(to=t["token"], title="SoulTracker - Resumen", body=body)
            for t in tokens
        ]
        try:
            client = PushClient()
            responses = client.publish_multiple(messages)
            # Store responses for audit
            for t, r in zip(tokens, responses):
                await db.push_logs.insert_one({
                    "manager": manager,
                    "token": t["token"],
                    "response": r.__dict__,
                    "period": period,
                    "created_at": datetime.utcnow(),
                })
        except PushServerError as e:
            logging.exception(f"Expo push error: {e}")


# ----------------------------------------------------------------------------
# Routes
# ----------------------------------------------------------------------------
@api_router.get("/")
async def root():
    return {"message": "SoulTracker API alive"}


@api_router.post("/status", response_model=StatusCheck)
async def create_status_check(input: StatusCheckCreate):
    status_obj = StatusCheck(**input.dict())
    await db.status_checks.insert_one(status_obj.dict())
    return status_obj


@api_router.get("/status", response_model=List[StatusCheck])
async def get_status_checks():
    status_checks = await db.status_checks.find().to_list(1000)
    return [StatusCheck(**status_check) for status_check in status_checks]


@api_router.post("/push/register")
async def register_push_token(data: PushTokenIn):
    if not data.token.startswith("ExponentPushToken"):
        raise HTTPException(status_code=400, detail="Invalid Expo push token")
    await db.push_tokens.update_one(
        {"token": data.token},
        {"$set": {"manager": data.manager, "updated_at": datetime.utcnow()}},
        upsert=True,
    )
    return {"ok": True}


@api_router.post("/upload-report")
async def upload_report(file: UploadFile = File(...)):
    try:
        if not file.filename or not (file.filename.endswith('.xlsx') or file.filename.endswith('.xlsm')):
            raise HTTPException(status_code=400, detail="Only .xlsx files are supported")

        content = await file.read()
        
        # Add validation for empty file
        if not content:
            raise HTTPException(status_code=400, detail="Empty file provided")
        
        wb = load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active

        headers = [c.value if c.value is not None else "" for c in next(ws.iter_rows(min_row=1, max_row=1))[0:ws.max_column]]
        header_map = normalize_header_map(headers)

        # Validate required headers - check for column variants
        required_checks = [
            (["Creator ID"], "Creator ID"),
            (["Creator's username"], "Creator's username"), 
            (["Joined time"], "Joined time"),
            (["Diamonds"], "Diamonds"),
            (["LIVE duration", "LIVE duration(h)"], "LIVE duration"),
            (["Valid go LIVE days", "Valid days(d)"], "Valid go LIVE days"),
        ]
        
        for variants, display_name in required_checks:
            found = any(variant.strip().lower() in header_map for variant in variants)
            if not found:
                raise HTTPException(status_code=400, detail=f"Missing required column: {display_name} (accepted variants: {', '.join(variants)})")

        processed = 0
        period = None

        # Iterate over rows
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {headers[i]: row[i] for i in range(len(headers))}
            if period is None:
                period = get_cell(row_dict, header_map, ["Data period"]) or "Unknown"

            creator_id = str(get_cell(row_dict, header_map, ["Creator ID"]))
            username = get_cell(row_dict, header_map, ["Creator's username"]) or ""
            manager = get_cell(row_dict, header_map, ["manager", "creator network manager"]) or "Unassigned"

            # Parse numbers
            diamonds = get_cell(row_dict, header_map, ["Diamonds"]) or 0
            try:
                diamonds = int(float(diamonds))
            except Exception:
                diamonds = 0

            live_duration = get_cell(row_dict, header_map, ["LIVE duration", "LIVE duration(h)"]) or 0
            try:
                live_duration = float(live_duration)
            except Exception:
                live_duration = 0.0

            valid_days = get_cell(row_dict, header_map, ["Valid go LIVE days", "Valid days(d)"]) or 0
            try:
                valid_days = int(float(valid_days))
            except Exception:
                valid_days = 0

            # Parse previous month data for growth analysis
            prev_diamonds = get_cell(row_dict, header_map, ["prev_diamonds", "diamonds last month"]) or 0
            try:
                prev_diamonds = int(float(prev_diamonds)) if prev_diamonds else 0
            except Exception:
                prev_diamonds = 0

            prev_live_duration = get_cell(row_dict, header_map, ["prev_live_duration_h", "live duration (hours) last month"]) or 0
            try:
                prev_live_duration = float(prev_live_duration) if prev_live_duration else 0
            except Exception:
                prev_live_duration = 0.0

            prev_valid_days = get_cell(row_dict, header_map, ["prev_valid_live_days", "valid go live days last month"]) or 0
            try:
                prev_valid_days = int(float(prev_valid_days)) if prev_valid_days else 0
            except Exception:
                prev_valid_days = 0

            joined_time = get_cell(row_dict, header_map, ["Joined time"]) or None
            if isinstance(joined_time, datetime):
                joined_dt = joined_time
            else:
                try:
                    joined_dt = datetime.fromisoformat(str(joined_time)) if joined_time else None
                except Exception:
                    joined_dt = None

            days_since_joining = get_cell(row_dict, header_map, ["Days since joining"]) or None
            if days_since_joining is None and joined_dt is not None:
                days_since_joining = (datetime.utcnow().date() - joined_dt.date()).days
            try:
                days_since_joining = int(days_since_joining) if days_since_joining is not None else None
            except Exception:
                days_since_joining = None

            goals_info = compute_goal_status(valid_days, live_duration)
            milestones = diamonds_milestones_near(diamonds)
            new_joiner = (days_since_joining is not None and days_since_joining < 90)

            alerts = []
            if goals_info.get("near_goal"):
                ng = goals_info["near_goal"]
                alerts.append({
                    "type": "goal",
                    "label": ng["label"],
                    "days_left": ng["days_left"],
                    "hours_left": ng["hours_left"],
                })
            for m in milestones:
                alerts.append({"type": "diamonds", "milestone": m["milestone"], "progress": m["progress"]})
            if new_joiner:
                alerts.append({"type": "new_joiner", "message": "Menos de 90 días desde que ingresó"})

            doc = {
                "creator_id": creator_id,
                "username": username,
                "manager": manager,
                "period": period,
                "joined_time": joined_dt,
                "days_since_joining": days_since_joining,
                "diamonds": diamonds,
                "live_duration_h": live_duration,
                "valid_live_days": valid_days,
                "prev_diamonds": prev_diamonds,
                "prev_live_duration_h": prev_live_duration,
                "prev_valid_live_days": prev_valid_days,
                "goals": goals_info,
                "milestones_near": milestones,
                "new_joiner": new_joiner,
                "alerts": alerts,
                "created_at": datetime.utcnow(),
            }
            await upsert_creator(doc)
            await db.reports.insert_one(doc)
            processed += 1

        # Fire and forget push notifications (do not block request)
        try:
            asyncio.create_task(push_summary_notifications(period or "Reporte"))
        except Exception:
            logging.exception("Failed to schedule push notifications task")

        return {"ok": True, "processed": processed, "period": period}
    
    except HTTPException:
        # Re-raise HTTP exceptions as they are already properly formatted
        raise
    except Exception as e:
        # Log the error and return a generic error response
        logging.exception(f"Error processing upload: {e}")
        raise HTTPException(status_code=500, detail="Internal server error while processing file")


@api_router.get("/creators")
async def list_creators(manager: Optional[str] = None, search: Optional[str] = None, alerts_only: Optional[bool] = False):
    q: Dict[str, Any] = {}
    if manager:
        q["manager"] = manager
    if alerts_only:
        q["alerts.0"] = {"$exists": True}
    cursor = db.creators.find(q).sort("username", 1)
    items = await cursor.to_list(2000)
    if search:
        s = search.lower()
        items = [i for i in items if s in (i.get("username", "").lower())]
    # Clean ObjectId
    for i in items:
        i["_id"] = str(i.get("_id"))
        if i.get("joined_time") and isinstance(i["joined_time"], datetime):
            i["joined_time"] = i["joined_time"].isoformat()
    return {"items": items}


@api_router.get("/alerts")
async def get_alerts(manager: Optional[str] = None):
    q: Dict[str, Any] = {}
    if manager:
        q["manager"] = manager
    q["alerts.0"] = {"$exists": True}
    items = await db.creators.find(q, {"username": 1, "alerts": 1}).to_list(2000)
    for i in items:
        i["_id"] = str(i.get("_id"))
    return {"items": items, "count": len(items)}


@api_router.get("/kpis")
async def get_kpis(manager: Optional[str] = None):
    q: Dict[str, Any] = {"manager": manager} if manager else {}
    total = await db.creators.count_documents(q)
    alerts = await db.creators.count_documents({**q, "alerts.0": {"$exists": True}})
    superstar = await db.creators.count_documents({**q, "goals.status.code": "SUPERSTAR"})
    return {
        "active_creators": total,
        "alerts_count": alerts,
        "superstars": superstar,
    }

# Include the router in the main app
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()