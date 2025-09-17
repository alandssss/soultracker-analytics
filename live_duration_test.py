#!/usr/bin/env python3
"""
LIVE DURATION CALCULATION DIAGNOSTIC TEST
Specific test for the reported issue with live_duration_h calculation
"""

import requests
import json
import os
from openpyxl import load_workbook
from datetime import datetime

# Get backend URL from frontend environment
def get_backend_url():
    try:
        with open('/app/frontend/.env', 'r') as f:
            for line in f:
                if line.startswith('EXPO_PUBLIC_BACKEND_URL='):
                    return line.split('=', 1)[1].strip()
    except Exception as e:
        print(f"Error reading frontend .env: {e}")
    return "https://nginx-stack-update.preview.emergentagent.com"

BASE_URL = get_backend_url()
API_BASE = f"{BASE_URL}/api"

print(f"🔍 LIVE DURATION DIAGNOSTIC TEST")
print(f"Testing backend at: {API_BASE}")
print("=" * 80)

def analyze_sample_xlsx():
    """Analyze the sample XLSX file to extract expected live_duration values"""
    print("\n=== STEP 1: ANALYZING SAMPLE XLSX FILE ===")
    
    try:
        wb = load_workbook('/app/sample_soultracker_data.xlsx', data_only=True)
        ws = wb.active
        
        # Get headers
        headers = [c.value if c.value is not None else '' for c in next(ws.iter_rows(min_row=1, max_row=1))[0:ws.max_column]]
        print(f"Headers found: {headers}")
        
        # Create header mapping (same logic as backend)
        header_map = {}
        for h in headers:
            header_map[h.strip().lower()] = h
        print(f"Header mapping: {header_map}")
        
        # Find LIVE duration column
        live_duration_variants = ["live duration", "live duration(h)"]
        live_duration_header = None
        for variant in live_duration_variants:
            if variant.lower() in header_map:
                live_duration_header = header_map[variant.lower()]
                break
        
        if not live_duration_header:
            print("❌ CRITICAL: LIVE duration column not found!")
            return None
        
        print(f"✅ LIVE duration column found: '{live_duration_header}'")
        
        # Extract data
        expected_data = []
        for row_num in range(2, ws.max_row + 1):
            row = list(ws.iter_rows(min_row=row_num, max_row=row_num, values_only=True))[0]
            row_dict = {headers[i]: row[i] for i in range(len(headers))}
            
            creator_id = str(row_dict.get("Creator ID", ""))
            username = row_dict.get("Creator's username", "")
            live_duration_raw = row_dict.get(live_duration_header)
            
            # Apply same conversion logic as backend
            try:
                live_duration_converted = float(live_duration_raw) if live_duration_raw is not None else 0.0
            except Exception:
                live_duration_converted = 0.0
            
            expected_data.append({
                "creator_id": creator_id,
                "username": username,
                "live_duration_raw": live_duration_raw,
                "live_duration_expected": live_duration_converted
            })
            
            print(f"Creator {creator_id} ({username}): Raw='{live_duration_raw}' ({type(live_duration_raw)}) -> Expected={live_duration_converted}")
        
        return expected_data
        
    except Exception as e:
        print(f"❌ Error analyzing XLSX: {e}")
        return None

def upload_sample_file():
    """Upload the sample XLSX file and return response"""
    print("\n=== STEP 2: UPLOADING SAMPLE XLSX FILE ===")
    
    try:
        with open('/app/sample_soultracker_data.xlsx', 'rb') as f:
            files = {'file': ('sample_soultracker_data.xlsx', f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
            response = requests.post(f"{API_BASE}/upload-report", files=files)
        
        print(f"Upload Status Code: {response.status_code}")
        print(f"Upload Response: {response.json()}")
        
        if response.status_code == 200:
            data = response.json()
            if data.get("ok") is True:
                print(f"✅ Upload successful - Processed: {data.get('processed', 0)}")
                return True
            else:
                print(f"❌ Upload failed - Response: {data}")
                return False
        else:
            print(f"❌ Upload failed - Status: {response.status_code}")
            return False
            
    except Exception as e:
        print(f"❌ Upload error: {e}")
        return False

def verify_database_values(expected_data):
    """Verify that live_duration_h values in database match expected values"""
    print("\n=== STEP 3: VERIFYING DATABASE VALUES ===")
    
    try:
        response = requests.get(f"{API_BASE}/creators")
        
        if response.status_code != 200:
            print(f"❌ Failed to get creators - Status: {response.status_code}")
            return False
        
        creators_data = response.json()
        creators_list = creators_data.get("items", [])
        
        print(f"Found {len(creators_list)} creators in database")
        
        # Create lookup by creator_id for comparison
        db_creators = {str(c.get("creator_id", "")): c for c in creators_list}
        
        print("\n--- LIVE DURATION COMPARISON ---")
        all_match = True
        
        for expected in expected_data:
            creator_id = expected["creator_id"]
            username = expected["username"]
            expected_value = expected["live_duration_expected"]
            
            if creator_id in db_creators:
                db_creator = db_creators[creator_id]
                actual_value = db_creator.get("live_duration_h")
                
                match = abs(float(actual_value or 0) - float(expected_value)) < 0.001  # Allow small floating point differences
                status = "✅ MATCH" if match else "❌ MISMATCH"
                
                print(f"{status} - {username} (ID: {creator_id})")
                print(f"  Expected: {expected_value}")
                print(f"  Actual:   {actual_value}")
                print(f"  Raw XLSX: {expected['live_duration_raw']}")
                
                if not match:
                    all_match = False
                    print(f"  ⚠️  DIFFERENCE: {abs(float(actual_value or 0) - float(expected_value))}")
            else:
                print(f"❌ MISSING - Creator {creator_id} ({username}) not found in database")
                all_match = False
        
        return all_match
        
    except Exception as e:
        print(f"❌ Database verification error: {e}")
        return False

def test_get_cell_function():
    """Test the get_cell function logic with various scenarios"""
    print("\n=== STEP 4: TESTING GET_CELL FUNCTION LOGIC ===")
    
    # Simulate the backend's get_cell function
    def get_cell(row_dict, header_map, key_variants, default=None):
        for kv in key_variants:
            lk = kv.lower()
            if lk in header_map:
                return row_dict.get(header_map[lk], default)
        return default
    
    # Test scenarios
    test_scenarios = [
        {
            "name": "Standard LIVE duration column",
            "headers": ["Creator ID", "LIVE duration", "Diamonds"],
            "row_data": {"Creator ID": "123", "LIVE duration": 45.5, "Diamonds": 1000},
            "variants": ["LIVE duration", "LIVE duration(h)"],
            "expected": 45.5
        },
        {
            "name": "LIVE duration(h) column",
            "headers": ["Creator ID", "LIVE duration(h)", "Diamonds"],
            "row_data": {"Creator ID": "123", "LIVE duration(h)": 65.2, "Diamonds": 1000},
            "variants": ["LIVE duration", "LIVE duration(h)"],
            "expected": 65.2
        },
        {
            "name": "Case insensitive matching",
            "headers": ["Creator ID", "live duration", "Diamonds"],
            "row_data": {"Creator ID": "123", "live duration": 25.8, "Diamonds": 1000},
            "variants": ["LIVE duration", "LIVE duration(h)"],
            "expected": 25.8
        },
        {
            "name": "Missing column",
            "headers": ["Creator ID", "Diamonds"],
            "row_data": {"Creator ID": "123", "Diamonds": 1000},
            "variants": ["LIVE duration", "LIVE duration(h)"],
            "expected": None
        }
    ]
    
    all_passed = True
    
    for scenario in test_scenarios:
        print(f"\nTesting: {scenario['name']}")
        
        # Create header mapping
        header_map = {}
        for h in scenario["headers"]:
            header_map[h.strip().lower()] = h
        
        # Test get_cell function
        result = get_cell(scenario["row_data"], header_map, scenario["variants"])
        
        if result == scenario["expected"]:
            print(f"✅ PASSED - Result: {result}")
        else:
            print(f"❌ FAILED - Expected: {scenario['expected']}, Got: {result}")
            all_passed = False
    
    return all_passed

def test_type_conversion():
    """Test type conversion scenarios"""
    print("\n=== STEP 5: TESTING TYPE CONVERSION ===")
    
    test_values = [
        {"input": 45.5, "expected": 45.5, "description": "Float value"},
        {"input": "65.2", "expected": 65.2, "description": "String float"},
        {"input": 25, "expected": 25.0, "description": "Integer value"},
        {"input": "30", "expected": 30.0, "description": "String integer"},
        {"input": None, "expected": 0.0, "description": "None value"},
        {"input": "", "expected": 0.0, "description": "Empty string"},
        {"input": "invalid", "expected": 0.0, "description": "Invalid string"},
        {"input": 0, "expected": 0.0, "description": "Zero value"},
    ]
    
    all_passed = True
    
    for test in test_values:
        print(f"\nTesting: {test['description']} - Input: {test['input']} ({type(test['input'])})")
        
        # Apply same conversion logic as backend
        try:
            result = float(test["input"]) if test["input"] is not None and test["input"] != "" else 0.0
        except Exception:
            result = 0.0
        
        if abs(result - test["expected"]) < 0.001:
            print(f"✅ PASSED - Result: {result}")
        else:
            print(f"❌ FAILED - Expected: {test['expected']}, Got: {result}")
            all_passed = False
    
    return all_passed

def run_comprehensive_live_duration_test():
    """Run the complete live duration diagnostic test"""
    print("🚀 STARTING COMPREHENSIVE LIVE DURATION DIAGNOSTIC")
    print("=" * 80)
    
    # Step 1: Analyze expected values from XLSX
    expected_data = analyze_sample_xlsx()
    if not expected_data:
        print("❌ CRITICAL FAILURE: Cannot analyze sample XLSX file")
        return False
    
    # Step 2: Upload file
    upload_success = upload_sample_file()
    if not upload_success:
        print("❌ CRITICAL FAILURE: File upload failed")
        return False
    
    # Step 3: Verify database values
    values_match = verify_database_values(expected_data)
    
    # Step 4: Test get_cell function logic
    get_cell_test = test_get_cell_function()
    
    # Step 5: Test type conversion
    conversion_test = test_type_conversion()
    
    # Summary
    print("\n" + "=" * 80)
    print("📊 LIVE DURATION DIAGNOSTIC SUMMARY")
    print("=" * 80)
    
    tests = [
        ("XLSX Analysis", expected_data is not None),
        ("File Upload", upload_success),
        ("Database Values Match", values_match),
        ("get_cell Function Logic", get_cell_test),
        ("Type Conversion Logic", conversion_test)
    ]
    
    all_passed = True
    for test_name, result in tests:
        status = "✅ PASSED" if result else "❌ FAILED"
        print(f"{test_name}: {status}")
        if not result:
            all_passed = False
    
    if all_passed:
        print("\n🎉 ALL TESTS PASSED - live_duration_h calculation is working correctly!")
    else:
        print("\n⚠️  ISSUES DETECTED - live_duration_h calculation has problems!")
    
    return all_passed

if __name__ == "__main__":
    run_comprehensive_live_duration_test()