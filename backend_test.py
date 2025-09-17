#!/usr/bin/env python3
"""
Backend API Testing for SoulTracker Analytics
Tests FastAPI endpoints according to review request specifications
"""

import requests
import json
import os
from io import BytesIO
import tempfile
from datetime import datetime

# Get backend URL from frontend environment
def get_backend_url():
    try:
        with open('/app/frontend/.env', 'r') as f:
            for line in f:
                if line.startswith('EXPO_PUBLIC_BACKEND_URL='):
                    return line.split('=', 1)[1].strip()
    except Exception as e:
        print(f"Error reading frontend .env: {e}")
    return "https://nginx-stack-update.preview.emergentagent.com"

BASE_URL = get_backend_url()
API_BASE = f"{BASE_URL}/api"

print(f"Testing backend at: {API_BASE}")

def test_health_endpoint():
    """Test 1: Verify health endpoint GET /api/"""
    print("\n=== Test 1: Health Check ===")
    try:
        response = requests.get(f"{API_BASE}/")
        print(f"Status Code: {response.status_code}")
        print(f"Response: {response.json()}")
        
        if response.status_code == 200:
            data = response.json()
            if data.get("message") == "SoulTracker API alive":
                print("✅ Health check PASSED")
                return True
            else:
                print(f"❌ Health check FAILED - Unexpected message: {data}")
                return False
        else:
            print(f"❌ Health check FAILED - Status: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ Health check FAILED - Error: {e}")
        return False

def create_sample_xlsx():
    """Create a sample XLSX file for testing"""
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        
        # Headers
        headers = [
            "Creator ID", "Creator's username", "Joined time", "Diamonds",
            "LIVE duration", "Valid go LIVE days", "manager", "Data period"
        ]
        for i, header in enumerate(headers, 1):
            ws.cell(row=1, column=i, value=header)
        
        # Sample data rows
        sample_data = [
            ["12345", "creator_alpha", "2024-01-15", 75000, 45.5, 8, "Manager_A", "09_2025"],
            ["12346", "creator_beta", "2024-02-20", 120000, 65.2, 22, "Manager_B", "09_2025"],
            ["12347", "creator_gamma", "2024-06-01", 35000, 25.8, 15, "Manager_A", "09_2025"],
            ["12348", "creator_delta", "2024-08-10", 280000, 85.4, 25, "Manager_C", "09_2025"],
        ]
        
        for row_idx, row_data in enumerate(sample_data, 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        # Save to temporary file
        temp_file = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        wb.save(temp_file.name)
        return temp_file.name
    except ImportError:
        print("❌ openpyxl not available, cannot create test XLSX")
        return None
    except Exception as e:
        print(f"❌ Error creating sample XLSX: {e}")
        return None

def test_upload_report():
    """Test 2: Upload XLSX report via POST /api/upload-report"""
    print("\n=== Test 2: Upload Report ===")
    
    xlsx_file = create_sample_xlsx()
    if not xlsx_file:
        print("❌ Cannot create test XLSX file")
        return False
    
    try:
        with open(xlsx_file, 'rb') as f:
            files = {'file': ('reporte_09_2025.xlsx', f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
            response = requests.post(f"{API_BASE}/upload-report", files=files)
        
        # Clean up temp file
        os.unlink(xlsx_file)
        
        print(f"Status Code: {response.status_code}")
        print(f"Response: {response.json()}")
        
        if response.status_code == 200:
            data = response.json()
            if (data.get("ok") is True and 
                data.get("processed", 0) > 0 and 
                "period" in data):
                print(f"✅ Upload PASSED - Processed: {data['processed']}, Period: {data['period']}")
                return True
            else:
                print(f"❌ Upload FAILED - Unexpected response structure: {data}")
                return False
        else:
            print(f"❌ Upload FAILED - Status: {response.status_code}")
            return False
            
    except Exception as e:
        print(f"❌ Upload FAILED - Error: {e}")
        if xlsx_file and os.path.exists(xlsx_file):
            os.unlink(xlsx_file)
        return False

def test_kpis_endpoint():
    """Test 3: Get KPIs via GET /api/kpis"""
    print("\n=== Test 3: KPIs Endpoint ===")
    try:
        response = requests.get(f"{API_BASE}/kpis")
        print(f"Status Code: {response.status_code}")
        print(f"Response: {response.json()}")
        
        if response.status_code == 200:
            data = response.json()
            required_fields = ["active_creators", "alerts_count", "superstars"]
            if all(field in data for field in required_fields):
                print("✅ KPIs endpoint PASSED")
                return True
            else:
                missing = [f for f in required_fields if f not in data]
                print(f"❌ KPIs endpoint FAILED - Missing fields: {missing}")
                return False
        else:
            print(f"❌ KPIs endpoint FAILED - Status: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ KPIs endpoint FAILED - Error: {e}")
        return False

def test_creators_endpoint():
    """Test 4: Get creators via GET /api/creators"""
    print("\n=== Test 4: Creators Endpoint ===")
    try:
        response = requests.get(f"{API_BASE}/creators")
        print(f"Status Code: {response.status_code}")
        print(f"Response structure: {type(response.json())}")
        
        if response.status_code == 200:
            data = response.json()
            if "items" in data and isinstance(data["items"], list):
                print(f"✅ Creators endpoint PASSED - Found {len(data['items'])} creators")
                
                # Check key fields in first creator if available
                if data["items"]:
                    creator = data["items"][0]
                    required_fields = ["username", "diamonds", "valid_live_days", "live_duration_h", "alerts"]
                    missing_fields = [f for f in required_fields if f not in creator]
                    if missing_fields:
                        print(f"⚠️  Missing fields in creator data: {missing_fields}")
                    else:
                        print("✅ Creator data structure is complete")
                        
                    # Check goals.status.code if goals exist
                    if "goals" in creator and "status" in creator["goals"] and "code" in creator["goals"]["status"]:
                        print(f"✅ Goals status found: {creator['goals']['status']['code']}")
                    else:
                        print("⚠️  Goals status structure incomplete")
                
                return True
            else:
                print(f"❌ Creators endpoint FAILED - Invalid response structure: {data}")
                return False
        else:
            print(f"❌ Creators endpoint FAILED - Status: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ Creators endpoint FAILED - Error: {e}")
        return False

def test_creators_alerts_filter():
    """Test 5: Get creators with alerts filter via GET /api/creators?alerts_only=true"""
    print("\n=== Test 5: Creators with Alerts Filter ===")
    try:
        # First get all creators
        all_response = requests.get(f"{API_BASE}/creators")
        alerts_response = requests.get(f"{API_BASE}/creators?alerts_only=true")
        
        print(f"All creators status: {all_response.status_code}")
        print(f"Alerts only status: {alerts_response.status_code}")
        
        if all_response.status_code == 200 and alerts_response.status_code == 200:
            all_data = all_response.json()
            alerts_data = alerts_response.json()
            
            all_count = len(all_data.get("items", []))
            alerts_count = len(alerts_data.get("items", []))
            
            print(f"All creators: {all_count}")
            print(f"Creators with alerts: {alerts_count}")
            
            if alerts_count <= all_count:
                print("✅ Alerts filter PASSED - Filtered count is valid")
                return True
            else:
                print(f"❌ Alerts filter FAILED - More alerts than total creators")
                return False
        else:
            print(f"❌ Alerts filter FAILED - Status codes: all={all_response.status_code}, alerts={alerts_response.status_code}")
            return False
    except Exception as e:
        print(f"❌ Alerts filter FAILED - Error: {e}")
        return False

def test_alerts_endpoint():
    """Test 6: Get alerts via GET /api/alerts"""
    print("\n=== Test 6: Alerts Endpoint ===")
    try:
        response = requests.get(f"{API_BASE}/alerts")
        print(f"Status Code: {response.status_code}")
        print(f"Response: {response.json()}")
        
        if response.status_code == 200:
            data = response.json()
            if "items" in data and "count" in data:
                items_length = len(data["items"])
                reported_count = data["count"]
                
                if items_length == reported_count:
                    print(f"✅ Alerts endpoint PASSED - Count matches items length: {reported_count}")
                    return True
                else:
                    print(f"❌ Alerts endpoint FAILED - Count mismatch: items={items_length}, count={reported_count}")
                    return False
            else:
                print(f"❌ Alerts endpoint FAILED - Missing items or count fields")
                return False
        else:
            print(f"❌ Alerts endpoint FAILED - Status: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ Alerts endpoint FAILED - Error: {e}")
        return False

def test_push_register():
    """Test 7: Register push token via POST /api/push/register"""
    print("\n=== Test 7: Push Token Registration ===")
    try:
        payload = {
            "manager": "QA",
            "token": "ExponentPushToken[xxxxxxxxxxxxxx]"
        }
        response = requests.post(f"{API_BASE}/push/register", json=payload)
        print(f"Status Code: {response.status_code}")
        print(f"Response: {response.json()}")
        
        if response.status_code == 200:
            data = response.json()
            if data.get("ok") is True:
                print("✅ Push token registration PASSED")
                return True
            else:
                print(f"❌ Push token registration FAILED - Unexpected response: {data}")
                return False
        else:
            print(f"❌ Push token registration FAILED - Status: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ Push token registration FAILED - Error: {e}")
        return False

def test_comprehensive_data_processing():
    """Test 8: COMPREHENSIVE DATA PROCESSING FLOW - Addresses user's specific issue"""
    print("\n=== Test 8: COMPREHENSIVE DATA PROCESSING FLOW ===")
    print("🔍 Testing complete data processing pipeline after XLSX upload")
    
    # Step 1: Upload fresh data
    print("\n--- Step 1: Upload Fresh XLSX Data ---")
    xlsx_file = create_sample_xlsx()
    if not xlsx_file:
        print("❌ Cannot create test XLSX file")
        return False
    
    upload_success = False
    processed_count = 0
    period = None
    
    try:
        with open(xlsx_file, 'rb') as f:
            files = {'file': ('comprehensive_test.xlsx', f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
            response = requests.post(f"{API_BASE}/upload-report", files=files)
        
        os.unlink(xlsx_file)
        
        if response.status_code == 200:
            data = response.json()
            if data.get("ok") is True and data.get("processed", 0) > 0:
                upload_success = True
                processed_count = data["processed"]
                period = data.get("period", "Unknown")
                print(f"✅ Upload successful - Processed: {processed_count}, Period: {period}")
            else:
                print(f"❌ Upload failed - Invalid response: {data}")
                return False
        else:
            print(f"❌ Upload failed - Status: {response.status_code}, Response: {response.text}")
            return False
            
    except Exception as e:
        print(f"❌ Upload failed - Error: {e}")
        if xlsx_file and os.path.exists(xlsx_file):
            os.unlink(xlsx_file)
        return False
    
    if not upload_success:
        return False
    
    # Step 2: Verify data appears in creators endpoint
    print("\n--- Step 2: Verify Data in Creators Endpoint ---")
    try:
        response = requests.get(f"{API_BASE}/creators")
        if response.status_code != 200:
            print(f"❌ Creators endpoint failed - Status: {response.status_code}")
            return False
        
        creators_data = response.json()
        creators_list = creators_data.get("items", [])
        
        if len(creators_list) < processed_count:
            print(f"❌ Data processing issue - Expected {processed_count} creators, found {len(creators_list)}")
            return False
        
        print(f"✅ Found {len(creators_list)} creators in database")
        
        # Verify data structure and processing
        sample_creator = creators_list[0] if creators_list else None
        if not sample_creator:
            print("❌ No creator data found")
            return False
        
        # Check required fields
        required_fields = ["username", "diamonds", "valid_live_days", "live_duration_h", "alerts", "goals"]
        missing_fields = [f for f in required_fields if f not in sample_creator]
        if missing_fields:
            print(f"❌ Missing required fields in creator data: {missing_fields}")
            return False
        
        # Check goals processing
        if "goals" not in sample_creator or "status" not in sample_creator["goals"]:
            print("❌ Goals processing failed - Missing goals.status")
            return False
        
        goals_status = sample_creator["goals"]["status"].get("code")
        if not goals_status:
            print("❌ Goals processing failed - Missing status code")
            return False
        
        print(f"✅ Goals processing verified - Status: {goals_status}")
        
        # Check alerts processing
        alerts = sample_creator.get("alerts", [])
        print(f"✅ Alerts processing verified - {len(alerts)} alerts found")
        
    except Exception as e:
        print(f"❌ Creators verification failed - Error: {e}")
        return False
    
    # Step 3: Verify KPIs calculation
    print("\n--- Step 3: Verify KPIs Calculation ---")
    try:
        response = requests.get(f"{API_BASE}/kpis")
        if response.status_code != 200:
            print(f"❌ KPIs endpoint failed - Status: {response.status_code}")
            return False
        
        kpis_data = response.json()
        
        # Verify KPIs match processed data
        active_creators = kpis_data.get("active_creators", 0)
        alerts_count = kpis_data.get("alerts_count", 0)
        superstars = kpis_data.get("superstars", 0)
        
        if active_creators < processed_count:
            print(f"❌ KPIs calculation issue - Expected at least {processed_count} active creators, got {active_creators}")
            return False
        
        print(f"✅ KPIs calculation verified:")
        print(f"   - Active creators: {active_creators}")
        print(f"   - Alerts count: {alerts_count}")
        print(f"   - Superstars: {superstars}")
        
    except Exception as e:
        print(f"❌ KPIs verification failed - Error: {e}")
        return False
    
    # Step 4: Verify alerts generation
    print("\n--- Step 4: Verify Alerts Generation ---")
    try:
        response = requests.get(f"{API_BASE}/alerts")
        if response.status_code != 200:
            print(f"❌ Alerts endpoint failed - Status: {response.status_code}")
            return False
        
        alerts_data = response.json()
        alerts_items = alerts_data.get("items", [])
        alerts_count_reported = alerts_data.get("count", 0)
        
        if len(alerts_items) != alerts_count_reported:
            print(f"❌ Alerts count mismatch - Items: {len(alerts_items)}, Count: {alerts_count_reported}")
            return False
        
        print(f"✅ Alerts generation verified - {alerts_count_reported} alerts generated")
        
        # Check alert structure
        if alerts_items:
            sample_alert = alerts_items[0]
            if "username" not in sample_alert or "alerts" not in sample_alert:
                print("❌ Alert structure invalid - Missing username or alerts")
                return False
            print(f"✅ Alert structure verified for user: {sample_alert['username']}")
        
    except Exception as e:
        print(f"❌ Alerts verification failed - Error: {e}")
        return False
    
    # Step 5: Verify data consistency across endpoints
    print("\n--- Step 5: Verify Data Consistency ---")
    try:
        # Get data from all endpoints
        creators_response = requests.get(f"{API_BASE}/creators")
        kpis_response = requests.get(f"{API_BASE}/kpis")
        alerts_response = requests.get(f"{API_BASE}/alerts")
        
        if not all(r.status_code == 200 for r in [creators_response, kpis_response, alerts_response]):
            print("❌ One or more endpoints failed during consistency check")
            return False
        
        creators_data = creators_response.json()
        kpis_data = kpis_response.json()
        alerts_data = alerts_response.json()
        
        # Count creators with alerts manually
        creators_with_alerts = sum(1 for c in creators_data["items"] if c.get("alerts", []))
        
        # Verify consistency
        kpis_alerts_count = kpis_data["alerts_count"]
        alerts_endpoint_count = alerts_data["count"]
        
        if not (creators_with_alerts == kpis_alerts_count == alerts_endpoint_count):
            print(f"❌ Data inconsistency detected:")
            print(f"   - Creators with alerts: {creators_with_alerts}")
            print(f"   - KPIs alerts count: {kpis_alerts_count}")
            print(f"   - Alerts endpoint count: {alerts_endpoint_count}")
            return False
        
        print(f"✅ Data consistency verified - All endpoints report {creators_with_alerts} creators with alerts")
        
    except Exception as e:
        print(f"❌ Data consistency check failed - Error: {e}")
        return False
    
    print("\n🎉 COMPREHENSIVE DATA PROCESSING TEST PASSED")
    print("✅ File upload successful (200 OK)")
    print("✅ Data processing and analysis working correctly")
    print("✅ All endpoints returning processed data")
    print("✅ KPIs calculation accurate")
    print("✅ Alerts generation functional")
    print("✅ Data consistency maintained across endpoints")
    
    return True

def run_all_tests():
    """Run all backend tests and return summary"""
    print("🚀 Starting SoulTracker Backend API Tests")
    print(f"Target URL: {API_BASE}")
    print("=" * 60)
    
    tests = [
        ("Health Check", test_health_endpoint),
        ("Upload Report", test_upload_report),
        ("KPIs Endpoint", test_kpis_endpoint),
        ("Creators Endpoint", test_creators_endpoint),
        ("Creators Alerts Filter", test_creators_alerts_filter),
        ("Alerts Endpoint", test_alerts_endpoint),
        ("Push Token Registration", test_push_register),
        ("COMPREHENSIVE Data Processing Flow", test_comprehensive_data_processing),
    ]
    
    results = {}
    for test_name, test_func in tests:
        results[test_name] = test_func()
    
    print("\n" + "=" * 60)
    print("📊 TEST SUMMARY")
    print("=" * 60)
    
    passed = 0
    failed = 0
    
    for test_name, result in results.items():
        status = "✅ PASSED" if result else "❌ FAILED"
        print(f"{test_name}: {status}")
        if result:
            passed += 1
        else:
            failed += 1
    
    print(f"\nTotal: {len(tests)} tests")
    print(f"Passed: {passed}")
    print(f"Failed: {failed}")
    
    return results

if __name__ == "__main__":
    run_all_tests()